
import pandas as pd

# Original Excel file
file_path = "data/Data Cleaning Start.xlsx"

# Load the correct worksheet with the proper header row
df = pd.read_excel(
    file_path,
    sheet_name="Sheet1 (2)",
    header=1
)

print("--- ORIGINAL DATA LOADED ---")
print("Rows:", len(df))
print("Columns:", len(df.columns))

print("\n--- COLUMN NAMES ---")
print(df.columns.tolist())

# Remove columns that are completely empty
df = df.dropna(axis=1, how="all")

print("\n--- EMPTY COLUMNS REMOVED ---")
print("Columns:", df.columns.tolist())
print("Shape:", df.shape)

# Remove completely duplicated rows
rows_before = len(df)
df = df.drop_duplicates().copy()
rows_removed = rows_before - len(df)

print("\n--- DUPLICATES REMOVED ---")
print("Rows before:", rows_before)
print("Rows after:", len(df))
print("Duplicates removed:", rows_removed)

# Clean client names by removing information inside parentheses
df["Client"] = (
    df["Client"]
    .str.replace(r"\s*\([^)]*\)", "", regex=True)
    .str.strip()
)

print("\n--- CLIENT NAMES CLEANED ---")
print(df["Client"].head(10).to_string(index=False))

# Clean contact names
df["Contact"] = (
    df["Contact"]
    .str.strip()
    .str.replace(r"\s+", " ", regex=True)
    .str.title()
)

print("\n--- CONTACT NAMES CLEANED ---")
print(df["Contact"].head(10).to_string(index=False))

# Split Department into Department and Region
split_columns = df["Department"].str.split(
    "_", n=1, expand=True
)

df["Department"] = split_columns[0].str.strip()
df["Region"] = split_columns[1].str.strip()

print("\n--- DEPARTMENT AND REGION SEPARATED ---")
print(df[["Department", "Region"]].head(10).to_string(index=False))

# Standardize payment method names
df["Payment"] = (
    df["Payment"]
    .astype("string")
    .str.strip()
    .str.title()
)

print("\n--- PAYMENT METHODS STANDARDIZED ---")
print(df["Payment"].value_counts(dropna=False).to_string())

# Check missing values in each column
print("\n--- MISSING VALUES CHECK ---")
print(df.isnull().sum().to_string())

# Check column data types
print("\n--- DATA TYPES CHECK ---")
print(df.dtypes.to_string())

# Check final working dataset dimensions
print("\n--- CURRENT DATASET SHAPE ---")
print(df.shape)

# Validate profit margin without changing the original values
valid_rows = (
    df["Revenue"].notna()
    & df["Profit"].notna()
    & (df["Revenue"] != 0)
    & df["Profit Margin"].notna()
)

calculated_margin = (
    df.loc[valid_rows, "Profit"]
    / df.loc[valid_rows, "Revenue"]
)

existing_margin = df.loc[valid_rows, "Profit Margin"]

mismatches = ~(
    (calculated_margin - existing_margin).abs() < 0.000001
)

print("\n--- PROFIT MARGIN VALIDATION ---")
print("Rows checked:", int(valid_rows.sum()))
print("Mismatches found:", int(mismatches.sum()))
print("Missing margins:", int(df["Profit Margin"].isna().sum()))

from pathlib import Path

# Create the output folder if it doesn't exist
output_folder = Path("output")
output_folder.mkdir(parents=True, exist_ok=True)

# Create a separate copy for Excel export
export_df = df.copy()

# Replace missing values with N/A in the exported copy only
export_df = export_df.fillna("N/A")

# Save the cleaned data to a new Excel file
output_path = output_folder / "cleaned_data.xlsx"
export_df.to_excel(output_path, index=False)

print("\n--- EXCEL FILE SAVED ---")
print("File:", output_path.resolve())
print("Rows:", len(export_df))
print("Columns:", len(export_df.columns))

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

# Open the cleaned Excel file
workbook = load_workbook(output_path)
worksheet = workbook.active

# Format the header row
for cell in worksheet[1]:
    cell.fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78"
    )
    cell.font = Font(
        bold=True,
        color="FFFFFF"
    )
    cell.alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

# Adjust column widths automatically
for column_cells in worksheet.columns:
    max_length = max(
        len(str(cell.value)) if cell.value is not None else 0
        for cell in column_cells
    )
    column_letter = column_cells[0].column_letter
    worksheet.column_dimensions[column_letter].width = max_length + 3

# Freeze the header row
worksheet.freeze_panes = "A2"

# Hide gridlines
worksheet.sheet_view.showGridLines = False

# Save formatting changes
workbook.save(output_path)

print("\n--- EXCEL FORMATTING COMPLETED ---")
print("Header styling: Dark blue with white bold text")
print("Column widths: Adjusted automatically")
print("Header row: Frozen")
print("Gridlines: Hidden")
print("Saved to:", output_path.resolve())

# Create an automated data-cleaning report
report_folder = Path("reports")
report_folder.mkdir(parents=True, exist_ok=True)

report_path = report_folder / "data_cleaning_report.txt"

report_content = f"""
DATA CLEANING & REPORTING AUTOMATION
===================================

1. DATASET SUMMARY
Original rows: 31
Final rows: {len(df)}
Columns after cleaning: {len(df.columns)}
Duplicate rows removed: {rows_removed}

2. CLEANING OPERATIONS
- Removed completely empty columns
- Removed duplicate rows
- Cleaned client names
- Standardized contact names
- Split Department and Region
- Standardized payment methods
- Preserved missing values during calculations
- Exported missing values as N/A in Excel
- Formatted the Excel worksheet

3. MISSING VALUES IN CLEANED DATA
{df.isnull().sum().to_string()}

4. PROFIT MARGIN VALIDATION
Rows checked: {int(valid_rows.sum())}
Mismatches found: {int(mismatches.sum())}
Missing profit margins: {int(df["Profit Margin"].isna().sum())}

5. OUTPUT FILE
{output_path.resolve()}
"""

report_path.write_text(report_content.strip(), encoding="utf-8")

print("\n--- DATA CLEANING REPORT CREATED ---")
print("Report:", report_path.resolve())
